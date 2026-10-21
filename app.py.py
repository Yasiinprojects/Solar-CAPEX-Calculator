import io
import math

import pandas as pd
import streamlit as st
from google import genai
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

st.set_page_config(page_title="Solar & BESS CapEx Engine", page_icon="⚡", layout="wide")
st.title("⚡ Solar & BESS Interconnection CapEx Engine")
st.caption(
    "Budget-stage estimate in NRM2-style BOQ format. All rates are indicative "
    "defaults: edit them to match your suppliers and labour costs."
)
st.markdown("---")

# ============================== INPUTS =======================================
sb = st.sidebar
sb.header("1. System & Grid")
kw = sb.number_input("Peak AC System Load (kW)", 50.0, 10000.0, 500.0, 50.0)
volt_label = sb.selectbox(
    "Grid Interconnection Voltage",
    ["0.415 kV (Low Voltage)", "11 kV (Medium Voltage)", "33 kV (Medium Voltage)"],
    index=1,
)
V = {"0.415": 415, "11 kV": 11000, "33 kV": 33000}[
    next(k for k in ("0.415", "11 kV", "33 kV") if volt_label.startswith(k))
]
is_lv = V == 415
pf = sb.slider("Power Factor (cos φ)", 0.80, 1.00, 0.90, 0.01)
include_tx = sb.checkbox(
    "Include step-up transformer", value=not is_lv, key=f"tx_{V}",
    help="LV connections normally need no step-up transformer.",
)

sb.header("2. Site & Installation")
length = sb.number_input("Cable Distance to Point of Connection (m)", 10.0, 2000.0, 150.0, 10.0)
ground = sb.selectbox(
    "Civil Trenching Condition",
    ["Soft Soil / Open Grass", "Mixed Ground / Hard Earth", "Existing Asphalt / Concrete Cutting"],
    index=1,
)
contingency = sb.slider("Commercial Contingency (%)", 0, 20, 10, 1) / 100.0
accuracy = sb.slider("Estimate accuracy range (±%)", 10, 40, 25, 5) / 100.0
if sb.button("Reset edited rates/quantities"):
    st.session_state.pop("boq_editor", None)

# ============================== ENGINEERING ==================================
current = kw * 1000 / (1.732 * V * pf)

TX_SIZES = [250, 500, 750, 1000, 1500, 2000, 2500, 5000]
req_kva = kw / pf * 1.20
n_tx = max(1, math.ceil(req_kva / TX_SIZES[-1]))
tx_kva = min(s for s in TX_SIZES if s >= req_kva / n_tx)
tx_qty = n_tx if include_tx else 0

CABLES = {
    415: ("4-Core {} mm² Cu XLPE/SWA/PVC (LV)", [(120, 250, 65), (240, 360, 110), (300, 410, 140)]),
    11000: ("3-Core {} mm² Cu XLPE Armoured (11 kV)", [(95, 270, 110), (185, 380, 145), (300, 490, 210)]),
    33000: ("3-Core {} mm² Cu XLPE Armoured (33 kV)", [(95, 250, 180), (185, 340, 230), (300, 430, 320)]),
}
DERATE = 0.85
tmpl, options = CABLES[V]
runs = 1
for size, amp, rate in options:
    if current <= amp * DERATE:
        break
else:
    size, amp, rate = options[-1]
    runs = math.ceil(current / (amp * DERATE))
cable_desc = tmpl.format(size) + (f", {runs} parallel runs" if runs > 1 else "")

RATINGS = [630, 1250, 2000, 2500, 3200, 4000]
need_a = current * 1.25
rating = next((r for r in RATINGS if r >= need_a), RATINGS[-1])
fault_ka = {415: 50, 11000: 25, 33000: 31.5}[V]
if is_lv:
    sw_name, sw_rate = "Main Switchboard (LV)", 8000 + 9 * rating
else:
    sw_name = "Ring Main Unit / MV Switchgear"
    sw_rate = (15000 + 8 * rating) * (1.8 if V == 33000 else 1.0)

trench_rate = {"Soft": 35.0, "Mixed": 55.0, "Existing": 95.0}[ground.split()[0]]
trench_rate *= 1 + 0.3 * (runs - 1)

tx_rate = tx_kva * 42.0
supply = tx_qty * tx_rate + sw_rate + length * runs * rate
labour = max(5000.0, 0.12 * supply)
protection = 7500.0 if is_lv else 14000.0
aux_cost = 0.0 if is_lv else 3500.0
earthing, testing = 4500.0, 9500.0
prelims = max(6000.0, 0.06 * (supply + length * trench_rate + labour + protection + aux_cost + earthing + testing))

# ============================== BOQ DATA =====================================
sig = (V, kw, ground)
if st.session_state.get("boq_sig") != sig:
    st.session_state.pop("boq_editor", None)
    st.session_state["boq_sig"] = sig

rows = [
    ("21.1.1", f"Step-Up Transformer ({tx_kva} kVA, Oil-Immersed)" + ("" if include_tx else " (excluded)"), float(tx_qty), "nr", tx_rate),
    ("21.2.4", f"{sw_name}, {rating} A, {fault_ka:g} kA assumed fault level", 1.0, "nr", float(sw_rate)),
    ("21.5.2", f"Interconnection Cabling: {cable_desc}", float(length * runs), "m", float(rate)),
    ("14.2.1", f"Civils: trench, bedding, backfill, warning tape ({ground})", float(length), "m", trench_rate),
    ("21.6.1", "Protection Relays & Revenue Metering", 1.0, "item", protection),
    ("21.6.3", "Substation Auxiliary DC Tripping Supply & Battery Charger" + (" (n/a at LV)" if is_lv else ""),
     0.0 if is_lv else 1.0, "item", 3500.0),
    ("21.7.1", "Earthing System & Bonding", 1.0, "item", earthing),
    ("21.9.3", "M&E Installation Labour (rigging, termination, glanding)", 1.0, "item", float(labour)),
    ("21.9.8", "Protection Testing, Relay Calibration & Commissioning", 1.0, "item", testing),
    ("02.1.2", "Site Preliminaries (fencing, permits, supervision, temp power)", 1.0, "item", float(prelims)),
]
df = pd.DataFrame(rows, columns=["Ref", "Item Description", "Qty", "Unit", "Unit Rate (USD)"])

assumptions = [
    ("Peak AC load (kW)", kw), ("Grid voltage", volt_label), ("Power factor", pf),
    ("Design current (A)", round(current, 1)), ("Cable derating factor", DERATE),
    ("Switchgear sizing", "1.25 × design current, next standard rating"),
    ("Assumed fault level (kA)", f"{fault_ka:g} (confirm with utility)"),
    ("Trenching condition", ground),
    ("Scope exclusions", "PV/BESS equipment, utility fees, studies, SCADA, land"),
    ("Status", "Budget-stage estimate only; verify before use in tenders"),
]

# ============================== TAB CREATION =================================
tab1, tab2, tab3 = st.tabs(["⚡ CapEx Calculator & BOQ", "📥 Export & Assumptions", "🤖 AI Co-Pilot"])

with tab1:
    if is_lv and (runs > 2 or kw > 1000):
        st.warning("Load is high for a 415 V connection. Check whether an MV connection is more realistic.")
    if need_a > RATINGS[-1]:
        st.warning(f"Required switchgear current ({need_a:,.0f} A) exceeds the largest standard rating listed.")
    if n_tx > 1 and include_tx:
        st.info(f"Load exceeds the largest single transformer: using {n_tx} × {tx_kva} kVA units.")

    st.subheader("⚙️ Editable Quantities & Unit Rates")
    st.caption("Edit Qty or Unit Rate. Edits reset when load, voltage or trenching condition changes; use the sidebar button to reset manually.")
    
    edited = st.data_editor(
        df, key="boq_editor", hide_index=True, use_container_width=True,
        disabled=["Ref", "Item Description", "Unit"],
        column_config={
            "Qty": st.column_config.NumberColumn(min_value=0.0, format="%.1f"),
            "Unit Rate (USD)": st.column_config.NumberColumn(min_value=0.0, format="$%.2f"),
        },
    )
    edited["Total (USD)"] = edited["Qty"] * edited["Unit Rate (USD)"]
    subtotal = edited["Total (USD)"].sum()
    cont_amt = subtotal * contingency
    grand = subtotal + cont_amt

    st.markdown("### Executive Dashboard")
    c1, c2, c3 = st.columns(3)
    c1.metric("Transformer", f"{n_tx} × {tx_kva} kVA" if include_tx else "None (LV)")
    c2.metric("Design Load Current", f"{current:,.1f} A")
    c3.metric("Grand Total CapEx", f"${grand:,.0f}")

    r1, r2 = st.columns(2)
    r1.metric(f"Low Estimate (-{accuracy:.0%})", f"${grand*(1-accuracy):,.0f}")
    r2.metric(f"High Estimate (+{accuracy:.0%})", f"${grand*(1+accuracy):,.0f}")

    st.info(
        f"Contingency: **${cont_amt:,.2f}** ({contingency:.0%}). "
        f"CapEx intensity: **${grand / kw:,.2f} per kW**. Covers the grid-connection scope only "
        "(excludes PV/BESS equipment, utility connection fees, studies and SCADA)."
    )

with tab2:
    st.subheader("📥 Excel Export")
    st.caption("Download your finalized estimate as a fully formatted professional workbook with live formulas and an assumptions audit trail.")

    def build_excel(df_boq, inputs, cont_pct, acc):
        wb = Workbook()
        ws = wb.active
        ws.title = "CapEx Estimate"
        head = ["Ref", "Item Description", "Qty", "Unit", "Unit Rate (USD)", "Total (USD)"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(horizontal="center")
        for i, r in enumerate(df_boq.itertuples(index=False), start=2):
            ws.append([r[0], r[1], r[2], r[3], r[4], f"=C{i}*E{i}"])
        last = len(df_boq) + 1
        s = last + 2
        pc, ca, g, pa, lo, hi = s + 1, s + 2, s + 3, s + 4, s + 5, s + 6
        ws[f"E{s}"], ws[f"F{s}"] = "Subtotal", f"=SUM(F2:F{last})"
        ws[f"E{pc}"], ws[f"F{pc}"] = "Contingency %", cont_pct
        ws[f"E{ca}"], ws[f"F{ca}"] = "Contingency amount", f"=F{s}*F{pc}"
        ws[f"E{g}"], ws[f"F{g}"] = "GRAND TOTAL", f"=F{s}+F{ca}"
        ws[f"E{pa}"], ws[f"F{pa}"] = "Accuracy range ±", acc
        ws[f"E{lo}"], ws[f"F{lo}"] = "Low estimate", f"=F{g}*(1-F{pa})"
        ws[f"E{hi}"], ws[f"F{hi}"] = "High estimate", f"=F{g}*(1+F{pa})"
        for row in range(2, last + 1):
            ws[f"C{row}"].number_format = '#,##0.0'
            ws[f"E{row}"].number_format = '#,##0.00'
            ws[f"F{row}"].number_format = '"$"#,##0.00'
        for col, w in zip("ABCDEF", (10, 70, 10, 8, 18, 18)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"

        wa = wb.create_sheet("Assumptions")
        wa.append(["Input / Assumption", "Value"])
        for k, v in inputs:
            wa.append([k, v])
        wa.column_dimensions["A"].width = 38
        wa.column_dimensions["B"].width = 70
        for c in wa[1]:
            c.font = Font(bold=True)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    st.download_button(
        "Download Professional Excel (.xlsx)",
        data=build_excel(edited[["Ref", "Item Description", "Qty", "Unit", "Unit Rate (USD)"]], assumptions, contingency, accuracy),
        file_name=f"CapEx_Estimate_{int(kw)}kW.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )

    st.markdown("---")
    st.subheader("📋 Registered Assumptions Audit Trail")
    assumptions_df = pd.DataFrame(assumptions, columns=["Parameter / Rule", "Current Value"])
    st.dataframe(assumptions_df, hide_index=True, use_container_width=True)

with tab3:
    st.subheader("🤖 AI Engineering Co-Pilot")
    st.caption("Ask questions about why specific line items were generated, what design assumptions mean, or what to verify with grid operators.")

    def get_secret(name):
        try:
            return st.secrets[name]
        except Exception:
            return None

    GEMINI_KEY = get_secret("GEMINI_API_KEY")
    GEMINI_MODEL = get_secret("GEMINI_MODEL")
    configured = bool(GEMINI_KEY and GEMINI_MODEL)

    if not configured:
        st.info("Assistant not configured yet: add both `GEMINI_API_KEY` and `GEMINI_MODEL` to your Streamlit secrets.")

    if "chat" not in st.session_state:
        st.session_state.chat = []

    def build_context():
        subtotal_str = f"USD {subtotal:,.2f}"
        grand_str = f"USD {grand:,.2f}"
        return f"""You are an expert engineering assistant inside a solar/BESS CapEx tool.
Load: {kw:,.0f} kW | Voltage: {volt_label}
Subtotal: {subtotal_str} | Grand Total: {grand_str}
Help the user understand their grid interconnection cost estimate, transformer sizing, cabling, and design assumptions."""

    for m in st.session_state.chat:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    if prompt := st.chat_input("Ask about this estimate…", disabled=not configured):
        st.session_state.chat.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)
        with st.chat_message("assistant"):
            try:
                client = genai.Client(api_key=GEMINI_KEY)
                chat_history = "\n".join([f"{m['role'].capitalize()}: {m['content']}" for m in st.session_state.chat[-4:]])
                full_prompt = f"{build_context()}\n\nRecent Conversation:\n{chat_history}"

                response_stream = client.models.generate_content_stream(
                    model=GEMINI_MODEL,
                    contents=full_prompt,
                )

                def token_stream():
                    for chunk in response_stream:
                        if chunk.text:
                            yield chunk.text

                reply = st.write_stream(token_stream())
                st.session_state.chat.append({"role": "assistant", "content": reply})
            except Exception as e:
                if "503" in str(e) or "UNAVAILABLE" in str(e):
                    st.warning("⚠️ Google's AI servers are experiencing temporary high traffic. Please wait 10 seconds and try sending again!")
                else:
                    st.error(f"The assistant is unavailable right now. Details: {type(e).__name__}")
                st.session_state.chat.pop()
