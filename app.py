import io
import math
import sqlite3
import pandas as pd
import requests
import streamlit as st
from google import genai
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

st.set_page_config(page_title="Solar & BESS CapEx Engine", page_icon="⚡", layout="wide")


# ============================== DATABASE INITIALIZATION ======================
def init_db():
    conn = sqlite3.connect("capex_projects.db", check_same_thread=False)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS projects (
            name TEXT PRIMARY KEY,
            region TEXT,
            currency TEXT,
            kw REAL,
            voltage TEXT,
            length REAL,
            ground TEXT,
            contingency REAL,
            accuracy REAL,
            include_tx INTEGER,
            bypass INTEGER,
            labor_idx REAL,
            trench_soft REAL,
            trench_mixed REAL,
            trench_exist REAL
        )
    ''')
    conn.commit()
    conn.close()

init_db()

def save_project_to_db(name, data):
    conn = sqlite3.connect("capex_projects.db", check_same_thread=False)
    c = conn.cursor()
    c.execute('''
        INSERT OR REPLACE INTO projects 
        (name, region, currency, kw, voltage, length, ground, contingency, accuracy, include_tx, bypass, labor_idx, trench_soft, trench_mixed, trench_exist)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (
        name, data["region"], data["currency"], data["kw"], data["voltage"], 
        data["length"], data["ground"], data["contingency"], data["accuracy"], 
        int(data["include_tx"]), int(data["bypass"]), data["labor_idx"], 
        data["trench_soft"], data["trench_mixed"], data["trench_exist"]
    ))
    conn.commit()
    conn.close()

def load_project_from_db(name):
    conn = sqlite3.connect("capex_projects.db", check_same_thread=False)
    c = conn.cursor()
    c.execute('SELECT * FROM projects WHERE name = ?', (name,))
    row = c.fetchone()
    conn.close()
    if row:
        return {
            "region": row[1], "currency": row[2], "kw": row[3], "voltage": row[4],
            "length": row[5], "ground": row[6], "contingency": row[7], "accuracy": row[8],
            "include_tx": bool(row[9]), "bypass": bool(row[10]), "labor_idx": row[11],
            "trench_soft": row[12], "trench_mixed": row[13], "trench_exist": row[14]
        }
    return None

def get_all_project_names():
    conn = sqlite3.connect("capex_projects.db", check_same_thread=False)
    c = conn.cursor()
    c.execute('SELECT name FROM projects')
    rows = c.fetchall()
    conn.close()
    return [r[0] for r in rows]

def delete_project_from_db(name):
    conn = sqlite3.connect("capex_projects.db", check_same_thread=False)
    c = conn.cursor()
    c.execute('DELETE FROM projects WHERE name = ?', (name,))
    conn.commit()
    conn.close()


# ============================== GLOBAL MASTER MARKET TRACKER ==================
COPPER_SYMBOL, COPPER_BASE, COPPER_SHARE = "HG=F", 4.00, 0.50
STEEL_SYMBOL, STEEL_BASE, STEEL_SHARE = "HRC=F", 1000.0, 0.30
YAHOO_CHART = "https://query1.finance.yahoo.com/v8/finance/chart/{}"


def fetch_yahoo_price(symbol):
    r = requests.get(YAHOO_CHART.format(symbol), params={"range": "1d", "interval": "1d"},
                     headers={"User-Agent": "Mozilla/5.0"}, timeout=5)
    r.raise_for_status()
    price = float(r.json()["chart"]["result"][0]["meta"]["regularMarketPrice"])
    return price if price > 0 else None


def to_index(price, base, share):
    return min(1.5, max(0.7, 1 + share * (price / base - 1)))


@st.cache_data(ttl=1800)  # Caches live market data for 30 minutes
def get_global_market_data():
    d = {
        "rates": {"USD": 1.0, "GBP": 0.78, "EUR": 0.92, "CNY": 7.20, "MYR": 4.70},
        "copper_index": 1.05, "steel_index": 1.02, "copper_price": None, "steel_price": None,
        "fx_live": False, "copper_live": False, "steel_live": False
    }
    try:
        r = requests.get("https://open.er-api.com/v6/latest/USD", timeout=5)
        if r.status_code == 200:
            data_rates = r.json().get("rates", {})
            for cur in ["USD", "GBP", "EUR", "CNY", "MYR"]:
                if cur in data_rates:
                    d["rates"][cur] = data_rates[cur]
            d["fx_live"] = True
    except Exception:
        pass

    for key, symbol, base, share in (("copper", COPPER_SYMBOL, COPPER_BASE, COPPER_SHARE),
                                     ("steel", STEEL_SYMBOL, STEEL_BASE, STEEL_SHARE)):
        try:
            price = fetch_yahoo_price(symbol)
            if price:
                d[f"{key}_price"], d[f"{key}_index"], d[f"{key}_live"] = price, to_index(price, base, share), True
        except Exception:
            pass

    flags = [d["fx_live"], d["copper_live"], d["steel_live"]]
    d["status"] = "GLOBAL MASTER CONNECTED" if all(flags) else "PARTIAL GLOBAL FEED" if any(flags) else "OFFLINE FALLBACK MODE"
    return d


live_market = get_global_market_data()


# ============================== EXPORT BUILDERS ===============================
def create_pdf_summary(df_boq, subtotal_val, cont_pct, cont_val, grand_val, kw_val, volt_label_val, region_val, curr_name, sym, mode_label):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleStyle", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#1F3864"), spaceAfter=6)
    sub = ParagraphStyle("SubtitleStyle", parent=styles["Normal"], fontSize=10, textColor=colors.HexColor("#555555"), spaceAfter=15)

    data = [["Ref", "Item Description", "Qty", "Unit", f"Rate ({sym})", f"Total ({sym})"]]
    data += [[str(r[0]), str(r[1]), f"{r[2]:,.1f}", str(r[3]), f"{sym}{r[4]:,.2f}", f"{sym}{r[2] * r[4]:,.2f}"]
             for r in df_boq.itertuples(index=False)]
    data += [["", "", "", "", label, f"{sym}{val:,.2f}"]
             for label, val in (("Subtotal", subtotal_val), (f"Contingency ({cont_pct:.0%})", cont_val), ("GRAND TOTAL", grand_val))]

    t = Table(data, colWidths=[45, 230, 45, 35, 80, 85])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"),
        ("ALIGN", (4, 0), (-1, -1), "RIGHT"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -4), 0.5, colors.HexColor("#DDDDDD")),
        ("FONTNAME", (4, -3), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (4, -1), (-1, -1), colors.HexColor("#EBF1F5")),
    ]))
    doc.build([
        Paragraph(f"⚡ Solar & BESS CapEx Executive Summary ({region_val})", title),
        Paragraph(f"Mode: {mode_label} | Currency: {curr_name} | Load: {kw_val:,.0f} kW | Voltage: {volt_label_val}", sub),
        Spacer(1, 10),
        t,
    ])
    return buffer.getvalue()


def build_excel(df_boq, inputs, cont_pct, acc, sym):
    wb = Workbook()
    ws = wb.active
    ws.title = "CapEx Estimate"
    ws.append(["Ref", "Item Description", "Qty", "Unit", f"Unit Rate ({sym})", f"Total ({sym})"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center")
    for i, r in enumerate(df_boq.itertuples(index=False), start=2):
        ws.append([r[0], r[1], r[2], r[3], r[4], f"=C{i}*E{i}"])

    last = len(df_boq) + 1
    s = last + 2
    pc, ca, g, pa, lo, hi = range(s + 1, s + 7)
    for row, label, val in (
        (s, "Subtotal", f"=SUM(F2:F{last})"), (pc, "Contingency %", cont_pct),
        (ca, "Contingency amount", f"=F{s}*F{pc}"), (g, "GRAND TOTAL", f"=F{s}+F{ca}"),
        (pa, "Accuracy range ±", acc), (lo, "Low estimate", f"=F{g}*(1-F{pa})"),
        (hi, "High estimate", f"=F{g}*(1+F{pa})"),
    ):
        ws[f"E{row}"], ws[f"F{row}"] = label, val
    money = f'"{sym}"#,##0.00'
    for row in range(2, last + 1):
        for col, fmt in (("C", "#,##0.0"), ("E", money), ("F", money)):
            ws[f"{col}{row}"].number_format = fmt
    for col, w in zip("ABCDEF", (10, 70, 10, 8, 18, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wa = wb.create_sheet("Assumptions")
    wa.append(["Input / Regional Audit Parameter", "Value"])
    for k, v in inputs:
        wa.append([k, str(v)])
    wa.column_dimensions["A"].width = 38
    wa.column_dimensions["B"].width = 70
    for c in wa[1]:
        c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


st.title("⚡ Solar & BESS Interconnection CapEx Engine")
st.caption("Global Master Tracker with Database Persistence, Multi-Currency Conversion & Manual UK Cost Controls.")
st.markdown("---")

# ============================== SIDEBAR CONTROLS =============================
sb = st.sidebar

# --- PROJECT SAVE / LOAD MANAGER ---
sb.header("💾 Project Vault (Database)")
saved_projects = get_all_project_names()
selected_saved_proj = sb.selectbox("Load Saved Project", ["-- Select Project --"] + saved_projects)

if selected_saved_proj != "-- Select Project --":
    if sb.button("📂 Load Project"):
        p_data = load_project_from_db(selected_saved_proj)
        if p_data:
            st.session_state["load_region"] = p_data["region"]
            st.session_state["load_currency"] = p_data["currency"]
            st.session_state["load_kw"] = p_data["kw"]
            st.session_state["load_voltage"] = p_data["voltage"]
            st.session_state["load_length"] = p_data["length"]
            st.session_state["load_ground"] = p_data["ground"]
            st.session_state["load_cont"] = p_data["contingency"]
            st.session_state["load_acc"] = p_data["accuracy"]
            st.session_state["load_tx"] = p_data["include_tx"]
            st.session_state["load_bypass"] = p_data["bypass"]
            st.session_state["load_labor"] = p_data["labor_idx"]
            st.session_state["load_t_soft"] = p_data["trench_soft"]
            st.session_state["load_t_mixed"] = p_data["trench_mixed"]
            st.session_state["load_t_exist"] = p_data["trench_exist"]
            st.success(f"Loaded '{selected_saved_proj}' successfully!")
            st.rerun()

    if sb.button("🗑️ Delete Project", type="secondary"):
        delete_project_from_db(selected_saved_proj)
        st.success(f"Deleted '{selected_saved_proj}'.")
        st.rerun()

sb.markdown("---")

sb.header("1. Presentation Currency")
currency_options = {
    "US Dollar ($)": {"code": "USD", "sym": "$"},
    "British Pound (£)": {"code": "GBP", "sym": "£"},
    "Euro (€)": {"code": "EUR", "sym": "€"},
    "Chinese Yuan (¥)": {"code": "CNY", "sym": "¥"},
    "Malaysian Ringgit (RM)": {"code": "MYR", "sym": "RM"}
}
default_curr_idx = 0
curr_keys = list(currency_options.keys())
if "load_currency" in st.session_state:
    for idx, label in enumerate(curr_keys):
        if currency_options[label]["code"] == st.session_state["load_currency"]:
            default_curr_idx = idx
            break

selected_curr_label = sb.selectbox("Display Estimate In", curr_keys, index=default_curr_idx)
curr_info = currency_options[selected_curr_label]
curr_code, sym = curr_info["code"], curr_info["sym"]

sb.header("2. Regional Market Selection")
regions = ["United Kingdom (UK)", "Global Baseline (USD)"]
default_reg_idx = 0
if "load_region" in st.session_state and st.session_state["load_region"] in regions:
    default_reg_idx = regions.index(st.session_state["load_region"])

selected_region = sb.selectbox("Project Region & Market Profile", regions, index=default_reg_idx)
is_uk = selected_region == "United Kingdom (UK)"

sb.header("3. Global Master Pricing Mode")
default_bypass = st.session_state.get("load_bypass", False)
bypass = sb.checkbox(
    "🔓 Manual Rate Mode (Bypass Global Tracker)", value=default_bypass,
    help="Disables live global commodity indexing and FX scaling to allow direct manual overrides.",
)

if bypass:
    copper_factor = steel_factor = 1.0
    fx_multiplier = 1.0
    mode_label, status_color, status_text = "Manual Rate Mode (Global Tracker Bypassed)", "#D97706", "MANUAL OVERRIDE ACTIVE"
else:
    copper_factor, steel_factor = live_market["copper_index"], live_market["steel_index"]
    usd_to_target_rate = live_market["rates"].get(curr_code, 1.0)
    fx_multiplier = usd_to_target_rate
    fallbacks = [n for n, k in (("FX", "fx_live"), ("copper", "copper_live"), ("steel", "steel_live")) if not live_market[k]]
    mode_label = ("Global Master Indexed (Live)" if not fallbacks else f"Global Master (defaults used for {', '.join(fallbacks)})")
    status_text = live_market["status"]
    status_color = {"GLOBAL MASTER CONNECTED": "#10B981", "PARTIAL GLOBAL FEED": "#F59E0B"}.get(status_text, "#EF4444")


def tile(label, value, color):
    return (f'<div style="background-color:#1E293B;padding:6px;border-radius:4px;">'
            f'<div style="color:#64748B;font-size:9px;">{label}</div>'
            f'<div style="color:{color};font-size:13px;font-weight:bold;">{value}</div></div>')


sb.markdown(
    f"""<div style="background-color:#0F172A;padding:14px;border-radius:8px;border:1px solid #1E293B;margin-bottom:15px;">
<div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:10px;">
<span style="color:#94A3B8;font-size:11px;font-weight:700;letter-spacing:0.5px;">GLOBAL MASTER TERMINAL</span>
<span style="background-color:{status_color};color:white;font-size:9px;font-weight:700;padding:2px 6px;border-radius:4px;">{status_text}</span></div>
<div style="display:grid;grid-template-columns:1fr 1fr;gap:8px;font-family:monospace;">
{tile(f"FX RATE ({curr_code})", f"{fx_multiplier:.3f}", "#F8FAFC")}{tile("COPPER INDEX", f"{copper_factor:.2f}x", "#38BDF8")}
{tile("STEEL INDEX", f"{steel_factor:.2f}x", "#38BDF8")}{tile("CURRENCY", curr_code, "#38BDF8")}
</div></div>""",
    unsafe_allow_html=True,
)

# ============================== UK MANUAL LOCAL COST CONTROLS =================
if is_uk:
    sb.header("4. UK Local In-Country Cost Controls")
    sb.caption("Enter proprietary UK labor rates, civil trenching norms, and compliance overheads (Base: GBP).")
    labor_idx = sb.number_input("UK M&E Labor Multiplier / Index", 0.80, 2.50, st.session_state.get("load_labor", 1.35), 0.05)
    trench_soft = sb.number_input("Trenching Rate - Soft Ground (£/m)", 20.0, 150.0, st.session_state.get("load_t_soft", 45.0), 5.0)
    trench_mixed = sb.number_input("Trenching Rate - Mixed Ground (£/m)", 40.0, 250.0, st.session_state.get("load_t_mixed", 75.0), 5.0)
    trench_exist = sb.number_input("Trenching Rate - Existing Asphalt (£/m)", 60.0, 400.0, st.session_state.get("load_t_exist", 125.0), 5.0)
    duty = sb.number_input("UKCA & Import Compliance Factor", 1.00, 1.20, 1.03, 0.01)
    voltages = ["0.400 kV (Low Voltage)", "11 kV (Medium Voltage)", "33 kV (Medium Voltage)"]
    gbp_to_usd = 1.0 / live_market["rates"].get("GBP", 0.78)
    trench_soft *= gbp_to_usd
    trench_mixed *= gbp_to_usd
    trench_exist *= gbp_to_usd
else:
    labor_idx = 1.00
    trench_soft, trench_mixed, trench_exist = 35.0, 55.0, 95.0
    duty = 1.00
    voltages = ["0.415 kV (Low Voltage)", "11 kV (Medium Voltage)", "33 kV (Medium Voltage)"]

sb.header("5. System & Grid Setup")
kw = sb.number_input("Peak AC System Load (kW)", 50.0, 10000.0, st.session_state.get("load_kw", 500.0), 50.0)

default_volt_idx = 1
if "load_voltage" in st.session_state and st.session_state["load_voltage"] in voltages:
    default_volt_idx = voltages.index(st.session_state["load_voltage"])

volt_label = sb.selectbox("Grid Interconnection Voltage", voltages, index=default_volt_idx)
V = next(v for k, v in (("0.400", 400), ("0.415", 415), ("11 kV", 11000), ("33 kV", 33000)) if volt_label.startswith(k))
is_lv = V in (400, 415)
pf = sb.slider("Power Factor (cos φ)", 0.80, 1.00, 0.90, 0.01)
include_tx = sb.checkbox("Include step-up transformer", value=st.session_state.get("load_tx", not is_lv), key=f"tx_{V}")

sb.header("6. Site & Installation")
length = sb.number_input("Cable Distance to Point of Connection (m)", 10.0, 2000.0, st.session_state.get("load_length", 150.0), 10.0)

grounds = ["Soft Soil / Open Grass", "Mixed Ground / Hard Earth", "Existing Asphalt / Concrete Cutting"]
default_ground_idx = 1
if "load_ground" in st.session_state and st.session_state["load_ground"] in grounds:
    default_ground_idx = grounds.index(st.session_state["load_ground"])

ground = sb.selectbox("Civil Trenching Condition", grounds, index=default_ground_idx)
contingency = sb.slider("Commercial Contingency (%)", 0, 20, int(st.session_state.get("load_cont", 0.10) * 100), 1) / 100.0
accuracy = sb.slider("Estimate accuracy range (±%)", 10, 40, int(st.session_state.get("load_acc", 0.25) * 100), 5) / 100.0
if sb.button("Reset edited rates/quantities"):
    st.session_state.pop("boq_editor", None)

# ============================== ENGINEERING ==================================
current = kw * 1000 / (1.732 * V * pf)

TX_SIZES = [250, 500, 750, 1000, 1500, 2000, 2500, 5000]
req_kva = kw / pf * 1.20
n_tx = max(1, math.ceil(req_kva / TX_SIZES[-1]))
tx_kva = min(s for s in TX_SIZES if s >= req_kva / n_tx)
tx_qty = n_tx if include_tx else 0

CABLES = {
    "LV": ("4-Core {} mm² Cu XLPE/SWA/PVC (LV)", [(120, 250, 65), (240, 360, 110), (300, 410, 140)]),
    11000: ("3-Core {} mm² Cu XLPE Armoured (11 kV)", [(95, 270, 110), (185, 380, 145), (300, 490, 210)]),
    33000: ("3-Core {} mm² Cu XLPE Armoured (33 kV)", [(95, 250, 180), (185, 340, 230), (300, 430, 320)]),
}
DERATE = 0.85
tmpl, options = CABLES["LV" if is_lv else V]
runs = 1
for size, amp, base_rate in options:
    if current <= amp * DERATE:
        break
else:
    size, amp, base_rate = options[-1]
    runs = math.ceil(current / (amp * DERATE))
rate = base_rate * copper_factor
cable_desc = tmpl.format(size) + (f", {runs} parallel runs" if runs > 1 else "")

RATINGS = [630, 1250, 2000, 2500, 3200, 4000]
need_a = current * 1.25
rating = next((r for r in RATINGS if r >= need_a), RATINGS[-1])
fault_ka = {400: 50, 415: 50, 11000: 25, 33000: 31.5}[V]

if is_lv:
    sw_name, sw_rate = "Main Switchboard (LV)", 8000 + 9 * rating
else:
    sw_name, sw_rate = "Ring Main Unit / MV Switchgear", (15000 + 8 * rating) * (1.8 if V == 33000 else 1.0)
sw_rate *= steel_factor * duty

base_trench_dict = {"Soft": trench_soft, "Mixed": trench_mixed, "Existing": trench_exist}
trench_rate = base_trench_dict[ground.split()[0]] * (1 + 0.3 * (runs - 1))
tx_rate = tx_kva * 42.0 * steel_factor * duty
supply = tx_qty * tx_rate + sw_rate + length * runs * rate
labour = max(5000.0, 0.12 * supply) * labor_idx
protection = (7500.0 if is_lv else 14000.0) * duty
aux_cost = (0.0 if is_lv else 3500.0) * duty
earthing, testing = 4500.0 * labor_idx, 9500.0 * labor_idx
prelims = max(6000.0, 0.06 * (supply + length * trench_rate + labour + protection + aux_cost + earthing + testing))

# ============================== BOQ DATA =====================================
sig = (selected_region, V, kw, ground, curr_code, bypass, labor_idx, trench_soft, trench_mixed, trench_exist)
if st.session_state.get("boq_sig") != sig:
    st.session_state.pop("boq_editor", None)
    st.session_state["boq_sig"] = sig

RATE, TOTAL = f"Unit Rate ({sym})", f"Total ({sym})"
rows = [
    ("21.1.1", f"Step-Up Transformer ({tx_kva} kVA, Oil-Immersed)" + ("" if include_tx else " (excluded)"), float(tx_qty), "nr", tx_rate),
    ("21.2.4", f"{sw_name}, {rating} A, {fault_ka:g} kA assumed fault level", 1.0, "nr", sw_rate),
    ("21.5.2", f"Interconnection Cabling: {cable_desc}", float(length * runs), "m", rate),
    ("14.2.1", f"Civils: trench, bedding, backfill, warning tape ({ground})", float(length), "m", trench_rate),
    ("21.6.1", "Protection Relays & Revenue Metering (Compliant)", 1.0, "item", protection),
    ("21.6.3", "Substation Auxiliary DC Tripping Supply & Battery Charger" + (" (n/a at LV)" if is_lv else ""),
     0.0 if is_lv else 1.0, "item", aux_cost),
    ("21.7.1", "Earthing System & Bonding", 1.0, "item", earthing),
    ("21.9.3", "M&E Installation Labour (rigging, termination, glanding)", 1.0, "item", labour),
    ("21.9.8", "Protection Testing, Relay Calibration & Commissioning", 1.0, "item", testing),
    ("02.1.2", "Site Preliminaries (fencing, permits, supervision, temp power)", 1.0, "item", prelims),
]
df = pd.DataFrame(rows, columns=["Ref", "Item Description", "Qty", "Unit", RATE])
df[RATE] = df[RATE].astype(float) * fx_multiplier

assumptions = [
    ("Project Region", selected_region), ("Engine Mode", mode_label), ("Presentation Currency", selected_curr_label),
    ("Live FX Rate Used", f"{fx_multiplier:.3f}"), ("Global Copper Index", f"{copper_factor:.2f}x"),
    ("Global Steel Index", f"{steel_factor:.2f}x"),
    ("Manual UK Labor Multiplier", f"{labor_idx}x" if is_uk else "N/A"),
    ("Peak AC load (kW)", kw), ("Grid voltage", volt_label),
    ("Power factor", pf), ("Design current (A)", round(current, 1)), ("Trenching condition", ground),
    ("Scope exclusions", "PV/BESS equipment, utility fees, studies, SCADA, land"),
    ("Status", f"Budget-stage estimate displayed in {selected_curr_label}"),
]

# ============================== TABS =========================================
tab1, tab2, tab3 = st.tabs(["⚡ CapEx Calculator & BOQ", "📥 Export & Assumptions", "🤖 AI Co-Pilot"])

with tab1:
    # --- SAVE CURRENT CONFIGURATION FORM ---
    with st.expander("💾 Save Current Project to Database"):
        proj_name_input = st.text_input("Project Name / ID", value="Site_Alpha_01")
        if st.button("Save Project"):
            if proj_name_input.strip():
                p_payload = {
                    "region": selected_region, "currency": curr_code, "kw": kw, "voltage": volt_label,
                    "length": length, "ground": ground, "contingency": contingency, "accuracy": accuracy,
                    "include_tx": include_tx, "bypass": bypass, "labor_idx": labor_idx if is_uk else 1.00,
                    "trench_soft": trench_soft, "trench_mixed": trench_mixed, "trench_exist": trench_exist
                }
                save_project_to_db(proj_name_input.strip(), p_payload)
                st.success(f"Project '{proj_name_input.strip()}' saved successfully!")
            else:
                st.warning("Please enter a valid project name.")

    if is_lv and (runs > 2 or kw > 1000):
        st.warning("Load is high for a Low Voltage connection. Check whether an MV connection is more realistic.")
    if need_a > RATINGS[-1]:
        st.warning(f"Required switchgear current ({need_a:,.0f} A) exceeds the largest standard rating listed.")
    if n_tx > 1 and include_tx:
        st.info(f"Load exceeds the largest single transformer: using {n_tx} × {tx_kva} kVA units.")

    st.subheader("⚙️ Editable Quantities & Unit Rates")
    st.caption(f"Rates displayed in **{selected_curr_label}**. Region: **{selected_region}** | Mode: **{mode_label}**.")
    edited = st.data_editor(
        df, key="boq_editor", hide_index=True, use_container_width=True,
        disabled=["Ref", "Item Description", "Unit"],
        column_config={
            "Qty": st.column_config.NumberColumn(min_value=0.0, format="%.1f"),
            RATE: st.column_config.NumberColumn(min_value=0.0, format=f"{sym}%.2f"),
        },
    )
    edited[TOTAL] = edited["Qty"] * edited[RATE]
    subtotal = edited[TOTAL].sum()
    cont_amt = subtotal * contingency
    grand = subtotal + cont_amt

    st.markdown("### Executive Dashboard")
    c1, c2, c3 = st.columns(3)
    c1.metric("Transformer", f"{n_tx} × {tx_kva} kVA" if include_tx else "None (LV)")
    c2.metric("Design Load Current", f"{current:,.1f} A")
    c3.metric("Grand Total CapEx", f"{sym}{grand:,.0f}")
    r1, r2 = st.columns(2)
    r1.metric(f"Low Estimate (-{accuracy:.0%})", f"{sym}{grand * (1 - accuracy):,.0f}")
    r2.metric(f"High Estimate (+{accuracy:.0%})", f"{sym}{grand * (1 + accuracy):,.0f}")

with tab2:
    st.subheader("📥 Export Reports")
    st.caption("Download your finalized estimate as an Excel workbook or a PDF executive summary.")
    boq_out = edited[["Ref", "Item Description", "Qty", "Unit", RATE]]
    col_ex, col_pdf = st.columns(2)
    col_ex.download_button(
        "📊 Download Professional Excel (.xlsx)",
        data=build_excel(boq_out, assumptions, contingency, accuracy, sym),
        file_name=f"CapEx_Estimate_{curr_code}_{int(kw)}kW.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary", use_container_width=True,
    )
    col_pdf.download_button(
        "📄 Download Executive Summary (.pdf)",
        data=create_pdf_summary(boq_out, subtotal, contingency, cont_amt, grand, kw, volt_label, selected_region, selected_curr_label, sym, mode_label),
        file_name=f"CapEx_Summary_{curr_code}_{int(kw)}kW.pdf",
        mime="application/pdf", type="secondary", use_container_width=True,
    )
    st.markdown("---")
    st.subheader("📋 Registered Regional & Currency Audit Trail")
    st.dataframe(pd.DataFrame(assumptions, columns=["Parameter / Regional Metric", "Current Value"]),
                 hide_index=True, use_container_width=True)

with tab3:
    st.subheader("🤖 AI Engineering Co-Pilot")
    st.caption("Ask questions about currency conversions, multi-currency reporting, or design specs.")

    def get_secret(name):
        try:
            return st.secrets[name]
        except Exception:
            return None

    GEMINI_KEY, GEMINI_MODEL = get_secret("GEMINI_API_KEY"), get_secret("GEMINI_MODEL")
    configured = bool(GEMINI_KEY and GEMINI_MODEL)
    if not configured:
        st.info("Assistant not configured yet: add both `GEMINI_API_KEY` and `GEMINI_MODEL` to your Streamlit secrets.")
    st.session_state.setdefault("chat", [])

    def build_context():
        return f"""You are an expert engineering assistant inside a solar/BESS CapEx tool.
Region: {selected_region} | Presentation Currency: {selected_curr_label} | Mode: {mode_label}
Load: {kw:,.0f} kW | Voltage: {volt_label} | Grand Total: {sym}{grand:,.2f}
Help the user understand their multi-currency cost estimate, global market tracking, and design assumptions."""

    for m in st.session_state.chat:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    if prompt := st.chat_input("Ask about this estimate…", disabled=not configured):
        st.session_state.chat.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)
        with st.chat_message("assistant"):
            try:
                history = "\n".join([f"{m['role'].capitalize()}: {m['content']}" for m in st.session_state.chat[-4:]])
                stream = genai.Client(api_key=GEMINI_KEY).models.generate_content_stream(
                    model=GEMINI_MODEL, contents=f"{build_context()}\n\nRecent Conversation:\n{history}")
                reply = st.write_stream(chunk.text for chunk in stream if chunk.text)
                st.session_state.chat.append({"role": "assistant", "content": reply})
            except Exception as e:
                if "503" in str(e) or "UNAVAILABLE" in str(e):
                    st.warning("⚠️ Google's AI servers are experiencing temporary high traffic. Please wait 10 seconds and try sending again!")
                else:
                    st.error(f"The assistant is unavailable right now. Details: {type(e).__name__}")
                st.session_state.chat.pop()
